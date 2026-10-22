"""
DOCX Document Parser (T-112).
Extracts headings, paragraphs, and tables from Microsoft Word (.docx) files.
"""

import io
from typing import List
from app.parsing.models import ExtractedContent, TextBlock

try:
    import docx
except ImportError:
    docx = None


class DOCXParser:
    @classmethod
    def parse(cls, file_bytes: bytes, filename: str = "document.docx") -> ExtractedContent:
        if not docx:
            # Fallback simple text decode
            text = file_bytes.decode("utf-8", errors="ignore")
            return ExtractedContent(
                full_text=text,
                text_blocks=[TextBlock(text=text, page_number=1, block_type="paragraph")],
                metadata={"format": "docx", "filename": filename}
            )

        doc = docx.Document(io.BytesIO(file_bytes))
        text_blocks: List[TextBlock] = []
        full_text_parts: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            b_type = "heading" if p.style and "heading" in p.style.name.lower() else "paragraph"
            full_text_parts.append(text)
            text_blocks.append(TextBlock(
                text=text,
                page_number=1,
                block_type=b_type
            ))

        # Extract table text
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    joined = " | ".join(row_texts)
                    full_text_parts.append(joined)
                    text_blocks.append(TextBlock(
                        text=joined,
                        page_number=1,
                        block_type="table_cell"
                    ))

        return ExtractedContent(
            full_text="\n\n".join(full_text_parts),
            text_blocks=text_blocks,
            images=[],
            metadata={"format": "docx", "filename": filename, "paragraph_count": len(text_blocks)}
        )
